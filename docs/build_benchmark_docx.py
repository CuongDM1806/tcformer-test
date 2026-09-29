from pathlib import Path
import sys

sys.path.insert(0, "/tmp/benchmark-docx-deps")

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html


ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "benchmark_bcic_iv_2a_loso.html"
OUTPUT = ROOT / "benchmark_bcic_iv_2a_loso.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_runs(paragraph, element):
    text = " ".join(element.text_content().split())
    links = element.xpath(".//a[@href]")
    if links:
        hrefs = [a.get("href") for a in links if a.get("href")]
        if hrefs:
            text += "\n" + " | ".join(hrefs)
    run = paragraph.add_run(text)
    if element.tag in {"b", "strong"}:
        run.bold = True
    return run


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "start" if edge == "left" else "end" if edge == "right" else edge
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "AAB7C4")
        borders.append(el)


tree = html.fromstring(INPUT.read_text(encoding="utf-8"))
body = tree.xpath("//body")[0]
doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.25)
section.right_margin = Cm(1.25)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9)
for name, size, color in (("Title", 20, "17365D"), ("Heading 1", 14, "1F4E78"), ("Heading 2", 11, "244F6F")):
    styles[name].font.name = "Arial"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

for element in body:
    tag = element.tag.lower() if isinstance(element.tag, str) else ""
    if tag == "h1":
        p = doc.add_paragraph(style="Title")
        p.add_run(" ".join(element.text_content().split()))
    elif tag == "h2":
        p = doc.add_paragraph(style="Heading 1")
        if "pagebreak" in (element.get("class") or ""):
            p.paragraph_format.page_break_before = True
        p.add_run(" ".join(element.text_content().split()))
    elif tag == "h3":
        doc.add_paragraph(" ".join(element.text_content().split()), style="Heading 2")
    elif tag == "p":
        p = doc.add_paragraph()
        add_runs(p, element)
    elif tag == "div":
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        shade(cell, "FFF2CC")
        set_cell_margins(cell, 110, 110, 110, 110)
        cell.text = " ".join(element.text_content().split())
    elif tag == "table":
        rows = element.xpath("./tr")
        if not rows:
            continue
        col_count = max(len(row.xpath("./th|./td")) for row in rows)
        table = doc.add_table(rows=len(rows), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        for row_idx, row in enumerate(rows):
            html_cells = row.xpath("./th|./td")
            row_class = row.get("class") or ""
            for col_idx, html_cell in enumerate(html_cells):
                cell = table.cell(row_idx, col_idx)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                text = " ".join(html_cell.text_content().split())
                urls = [a.get("href") for a in html_cell.xpath(".//a[@href]") if a.get("href")]
                if urls:
                    text += "\n" + "\n".join(urls)
                cell.text = text
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(7.4 if col_count >= 8 else 8.2)
                if row_idx == 0 or html_cell.tag == "th":
                    shade(cell, "1F4E78")
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                elif "good" in row_class:
                    shade(cell, "E2F0D9")
                elif "warn" in row_class:
                    shade(cell, "FFF2CC")
                elif "unclear" in row_class:
                    shade(cell, "F4CCCC")
        doc.add_paragraph()
    elif tag in {"ol", "ul"}:
        style = "List Number" if tag == "ol" else "List Bullet"
        for li in element.xpath("./li"):
            p = doc.add_paragraph(style=style)
            add_runs(p, li)

doc.core_properties.title = "Benchmark BCIC-IV-2a Cross-subject / LOSO"
doc.core_properties.subject = "Protocol-aware benchmark for HADA/TCFormer"
doc.core_properties.author = "Benchmark research notes"
doc.save(OUTPUT)
print(OUTPUT)
